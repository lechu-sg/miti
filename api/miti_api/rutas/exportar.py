"""Exportación de balance, movimientos y liquidación a Excel y PDF (§3.10 y §7.5 de DEFINICION.md)."""

import io
import uuid
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Response, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from .. import cripto
from ..db import sesion
from ..modelos import Caja, Campana, Gasto, Integrante, Movimiento, Producto, Usuario, Venta
from ..seguridad import Contexto, contexto_activo

ruteador = APIRouter(tags=["exportación"])


def _formatear_plata(centavos: int) -> str:
    pesos = centavos / 100.0
    return f"${pesos:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


@ruteador.get("/campanas/{campana_id}/exportar/excel")
async def exportar_excel(
    ctx: Contexto = Depends(contexto_activo),
    s: AsyncSession = Depends(sesion),
) -> Response:
    """Genera y descarga el libro Excel con todas las hojas de la campaña."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    # Cargar datos completos
    camp = ctx.campana
    integrantes = (
        await s.execute(
            select(Integrante)
            .options(selectinload(Integrante.usuario))
            .where(Integrante.campana_id == camp.id)
            .order_by(Integrante.alta.asc())
        )
    ).scalars().all()

    cajas = (
        await s.execute(
            select(Caja, Usuario.nombre)
            .join(Usuario, Caja.titular_id == Usuario.id)
            .where(Caja.campana_id == camp.id)
        )
    ).all()

    movs = (
        await s.execute(
            select(Movimiento)
            .where(Movimiento.campana_id == camp.id)
            .order_by(Movimiento.creado.asc())
        )
    ).scalars().all()

    ventas = (
        await s.execute(
            select(Venta)
            .options(
                selectinload(Venta.comprador),
                selectinload(Venta.vendedor),
                selectinload(Venta.items),
            )
            .where(Venta.campana_id == camp.id)
            .order_by(Venta.creada.asc())
        )
    ).scalars().all()

    gastos = (
        await s.execute(
            select(Gasto)
            .options(selectinload(Gasto.movimiento), selectinload(Gasto.caja))
            .where(Gasto.campana_id == camp.id)
            .order_by(Gasto.creado.asc())
        )
    ).scalars().all()

    wb = openpyxl.Workbook()

    # Estilos comunes
    fuente_titulo = Font(name="Arial", size=14, bold=True, color="1A1A1A")
    fuente_cabecera = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fuente_datos = Font(name="Arial", size=10)
    fill_cabecera = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    fill_sub = PatternFill(start_color="ECEFF1", end_color="ECEFF1", fill_type="solid")
    borde_fino = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # 1. Hoja Resumen
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.append(["MITI — BALANCE DE CAMPAÑA"])
    ws_resumen["A1"].font = fuente_titulo
    ws_resumen.append([])
    ws_resumen.append(["Campaña:", camp.nombre])
    ws_resumen.append(["Tipo:", "Rifa de números" if camp.tipo == "rifa" else "Venta de productos"])
    ws_resumen.append(["Estado:", camp.estado.upper()])
    ws_resumen.append(["Moneda:", camp.moneda])
    ws_resumen.append(["Fecha de reporte:", datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws_resumen.append([])

    # Calcular totales
    total_vendido = sum(v.importe for v in ventas if v.estado == "confirmada")
    cobros_conf = [m for m in movs if m.tipo == "cobro" and m.estado == "confirmado"]
    total_cobrado = sum(m.importe for m in cobros_conf)
    total_gastos = sum(g.movimiento.importe for g in gastos if g.movimiento and g.movimiento.estado == "confirmado")
    neto = max(0, total_cobrado - total_gastos)

    ws_resumen.append(["MÉTRICA", "VALOR"])
    ws_resumen.append(["Total Vendido", _formatear_plata(total_vendido)])
    ws_resumen.append(["Total Cobrado (en cajas)", _formatear_plata(total_cobrado)])
    ws_resumen.append(["Gastos Totales Aprobados", _formatear_plata(total_gastos)])
    ws_resumen.append(["Neto Cobrado a Repartir", _formatear_plata(neto)])
    ws_resumen.append(["Integrantes activos", len([i for i in integrantes if i.estado == "activo"])])

    for col in ("A", "B"):
        ws_resumen.column_dimensions[col].width = 28

    # 2. Hoja Cajas
    ws_cajas = wb.create_sheet(title="Cajas")
    ws_cajas.append(["CAJA", "TIPO", "TITULAR", "INGRESOS", "EGRESOS", "SALDO DISPONIBLE"])
    for cell in ws_cajas[1]:
        cell.font = fuente_cabecera
        cell.fill = fill_cabecera

    for c, titular_nombre in cajas:
        c_ing = sum(m.importe for m in movs if m.caja_destino == c.id and m.estado == "confirmado")
        c_egr = sum(m.importe for m in movs if m.caja_origen == c.id and m.estado == "confirmado")
        ws_cajas.append([
            str(c.id)[:8],
            c.tipo.capitalize(),
            titular_nombre,
            _formatear_plata(c_ing),
            _formatear_plata(c_egr),
            _formatear_plata(c_ing - c_egr),
        ])

    for col in ("A", "B", "C", "D", "E", "F"):
        ws_cajas.column_dimensions[col].width = 20

    # 3. Hoja Ventas
    ws_ventas = wb.create_sheet(title="Ventas")
    ws_ventas.append(["CÓDIGO", "FECHA", "VENDEDOR", "COMPRADOR", "TELÉFONO", "DETALLE", "TOTAL", "ESTADO"])
    for cell in ws_ventas[1]:
        cell.font = fuente_cabecera
        cell.fill = fill_cabecera

    for v in ventas:
        comp_nom = "Anónimo"
        comp_tel = "-"
        if v.comprador:
            comp_nom = cripto.descifrar(v.comprador.nombre_cifrado)
            if ctx.es_admin or v.vendedor_id == ctx.usuario.id:
                comp_tel = cripto.descifrar(v.comprador.telefono_cifrado)
            else:
                comp_tel = "Oculto (privacidad)"

        detalle_items = []
        if camp.tipo == "rifa":
            detalle_items = [str(it.numero) for it in v.items if it.numero is not None]
            det_str = "Números: " + ", ".join(detalle_items)
        else:
            detalle_items = [f"{it.cantidad}x {it.producto.nombre if it.producto else 'Item'}" for it in v.items]
            det_str = ", ".join(detalle_items)

        ws_ventas.append([
            v.codigo_corto,
            v.creada.strftime("%d/%m/%Y %H:%M") if v.creada else "-",
            v.vendedor.nombre if v.vendedor else "-",
            comp_nom,
            comp_tel,
            det_str,
            _formatear_plata(v.importe),
            v.estado.upper(),
        ])

    for col in ("A", "B", "C", "D", "E", "F", "G", "H"):
        ws_ventas.column_dimensions[col].width = 20
    ws_ventas.column_dimensions["F"].width = 35

    # 4. Hoja Gastos
    ws_gastos = wb.create_sheet(title="Gastos")
    ws_gastos.append(["FECHA", "DESCRIPCIÓN", "ORIGEN", "CAJA", "IMPORTE", "ESTADO"])
    for cell in ws_gastos[1]:
        cell.font = fuente_cabecera
        cell.fill = fill_cabecera

    for g in gastos:
        ws_gastos.append([
            g.creado.strftime("%d/%m/%Y %H:%M") if g.creado else "-",
            g.descripcion,
            "Bolsillo" if g.origen == "bolsillo" else "Caja",
            g.caja.tipo.capitalize() if g.caja else "-",
            _formatear_plata(g.movimiento.importe if g.movimiento else 0),
            g.movimiento.estado.upper() if g.movimiento else "-",
        ])

    for col in ("A", "B", "C", "D", "E", "F"):
        ws_gastos.column_dimensions[col].width = 22
    ws_gastos.column_dimensions["B"].width = 35

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()

    filename = f"miti_balance_{camp.nombre.lower().replace(' ', '_')}.xlsx"
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@ruteador.get("/campanas/{campana_id}/exportar/pdf")
async def exportar_pdf(
    ctx: Contexto = Depends(contexto_activo),
    s: AsyncSession = Depends(sesion),
) -> Response:
    """Genera y descarga el balance oficial de la campaña en PDF."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    camp = ctx.campana
    movs = (
        await s.execute(
            select(Movimiento)
            .where(Movimiento.campana_id == camp.id)
            .order_by(Movimiento.creado.asc())
        )
    ).scalars().all()

    ventas = (
        await s.execute(
            select(Venta)
            .where(Venta.campana_id == camp.id, Venta.estado == "confirmada")
        )
    ).scalars().all()

    gastos = (
        await s.execute(
            select(Gasto)
            .options(selectinload(Gasto.movimiento), selectinload(Gasto.caja))
            .where(Gasto.campana_id == camp.id)
        )
    ).scalars().all()

    integrantes = (
        await s.execute(
            select(Integrante)
            .options(selectinload(Integrante.usuario))
            .where(Integrante.campana_id == camp.id, Integrante.estado.in_(("activo", "retirado")))
        )
    ).scalars().all()

    total_vendido = sum(v.importe for v in ventas)
    total_cobrado = sum(m.importe for m in movs if m.tipo == "cobro" and m.estado == "confirmado")
    total_gastos = sum(g.movimiento.importe for g in gastos if g.movimiento and g.movimiento.estado == "confirmado")
    neto = max(0, total_cobrado - total_gastos)
    n = max(1, len(integrantes))
    cuota = neto // n

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    story = []
    styles = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle(
        "TituloMiti",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1A1A1A"),
    )
    estilo_sub = ParagraphStyle(
        "SubMiti",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#666666"),
    )
    estilo_seccion = ParagraphStyle(
        "SeccionMiti",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#C0392B"),
    )
    estilo_celda = ParagraphStyle(
        "CeldaMiti",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
    )
    estilo_celda_negrita = ParagraphStyle(
        "CeldaNegritaMiti",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
    )

    story.append(Paragraph("MITI — RENDICIÓN Y BALANCE GENERAL", estilo_titulo))
    story.append(Paragraph(f"Campaña: <b>{camp.nombre}</b> · Estado: <b>{camp.estado.upper()}</b>", estilo_sub))
    story.append(Paragraph(f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')}", estilo_sub))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1A1A1A"), spaceAfter=14))

    # Resumen general
    story.append(Paragraph("1. RESUMEN CONTABLE", estilo_seccion))
    story.append(Spacer(1, 6))
    datos_resumen = [
        [Paragraph("Concepto", estilo_celda_negrita), Paragraph("Importe", estilo_celda_negrita)],
        [Paragraph("Total Vendido", estilo_celda), Paragraph(_formatear_plata(total_vendido), estilo_celda)],
        [Paragraph("Total Cobrado en Cajas", estilo_celda), Paragraph(_formatear_plata(total_cobrado), estilo_celda)],
        [Paragraph("Gastos Aprobados", estilo_celda), Paragraph(_formatear_plata(total_gastos), estilo_celda)],
        [Paragraph("Neto a Repartir", estilo_celda_negrita), Paragraph(_formatear_plata(neto), estilo_celda_negrita)],
        [Paragraph(f"Cuota Parte Individual ({n} integrantes)", estilo_celda_negrita), Paragraph(_formatear_plata(cuota), estilo_celda_negrita)],
    ]
    t_resumen = Table(datos_resumen, colWidths=[320, 200])
    t_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (1, 0), colors.HexColor("#F0F0F0")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t_resumen)
    story.append(Spacer(1, 16))

    # Integrantes
    story.append(Paragraph("2. NÓMINA DE INTEGRANTES", estilo_seccion))
    story.append(Spacer(1, 6))
    datos_integrantes = [
        [Paragraph("Nombre", estilo_celda_negrita), Paragraph("Rol", estilo_celda_negrita), Paragraph("Estado", estilo_celda_negrita), Paragraph("Firma", estilo_celda_negrita)]
    ]
    for i in integrantes:
        datos_integrantes.append([
            Paragraph(i.usuario.nombre, estilo_celda),
            Paragraph(i.rol.capitalize(), estilo_celda),
            Paragraph(i.estado.capitalize(), estilo_celda),
            Paragraph("________________________", estilo_celda),
        ])
    t_integrantes = Table(datos_integrantes, colWidths=[180, 80, 80, 180])
    t_integrantes.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0F0F0")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(t_integrantes)
    story.append(Spacer(1, 16))

    # Gastos
    story.append(Paragraph("3. GASTOS REGISTRADOS", estilo_seccion))
    story.append(Spacer(1, 6))
    datos_gastos = [
        [Paragraph("Descripción", estilo_celda_negrita), Paragraph("Origen", estilo_celda_negrita), Paragraph("Importe", estilo_celda_negrita), Paragraph("Estado", estilo_celda_negrita)]
    ]
    if gastos:
        for g in gastos:
            datos_gastos.append([
                Paragraph(g.descripcion, estilo_celda),
                Paragraph("Bolsillo" if g.origen == "bolsillo" else "Caja", estilo_celda),
                Paragraph(_formatear_plata(g.movimiento.importe if g.movimiento else 0), estilo_celda),
                Paragraph(g.movimiento.estado.upper() if g.movimiento else "-", estilo_celda),
            ])
    else:
        datos_gastos.append([Paragraph("Sin gastos registrados", estilo_celda), Paragraph("-", estilo_celda), Paragraph("$0,00", estilo_celda), Paragraph("-", estilo_celda)])

    t_gastos = Table(datos_gastos, colWidths=[240, 90, 100, 90])
    t_gastos.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0F0F0")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t_gastos)

    doc.build(story)
    pdf_bytes = buf.getvalue()

    filename = f"miti_balance_{camp.nombre.lower().replace(' ', '_')}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
